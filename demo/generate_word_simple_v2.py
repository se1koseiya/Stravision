"""
精简版：生成 Word 文档 - 只保留核心内容
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_word_document():
    doc = Document()
    
    # 设置中文字体
    doc.styles['Normal'].font.name = u'微软雅黑'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'微软雅黑')
    
    # 标题
    heading = doc.add_heading('提示词工程六步演进', level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 副标题
    subtitle = doc.add_paragraph('智慧农业平台 AI 助手开发流程')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()
    
    # 版本 1
    doc.add_heading('版本 1：基础问答', level=1)
    doc.add_paragraph('需求：回答用户关于大棚种植的基本问题')
    code1 = '''export const BASIC_PROMPT = `
你是一个智慧农业助手的 AI。
请回答用户关于大棚种植的问题。
`;'''
    doc.add_paragraph(code1, style='No Spacing')
    
    p1 = doc.add_paragraph()
    p1.add_run('问题：').bold = True
    doc.add_paragraph('无法获取实时数据、建议缺乏针对性、无法执行操作', style='List Bullet')
    
    # 版本 2
    doc.add_heading('版本 2：数据注入', level=1)
    doc.add_paragraph('需求：让 AI 知道当前大棚的实际情况')
    code2 = '''export const DATA_ENHANCED_PROMPT = `
你是一个智慧农业助手的 AI。

## 当前设备数据（实时注入）
- 室内温度：32.5°C
- 室内湿度：55%
- 光照强度：12000 lux
- CO₂浓度：380 ppm
- 土壤温度：28.2°C
- 土壤湿度：35%

请基于真实数据回答，严禁编造数据。
`;'''
    doc.add_paragraph(code2, style='No Spacing')
    
    doc.add_paragraph('效果：可以回答实时数据问题，建议更有针对性')
    
    # 版本 3
    doc.add_heading('版本 3：诊断逻辑', level=1)
    doc.add_paragraph('需求：让 AI 判断数据是否正常，给出专业建议')
    doc.add_paragraph('添加农业专家知识库：')
    params = [
        '温度：18-28°C（白天）',
        '湿度：60-75%',
        '光照：15000-30000 lux',
        'CO₂：400-800 ppm'
    ]
    for param in params:
        doc.add_paragraph(param, style='List Bullet')
    
    doc.add_paragraph('效果：专业的数据对比和问题识别')
    
    # 版本 4
    doc.add_heading('版本 4：任务规划', level=1)
    doc.add_paragraph('需求：制定执行计划，告诉 AI 先做什么后做什么')
    task_list = '''<TASK_LIST>
{
  "title": "大棚环境优化计划",
  "tasks": [
    {"id": "1", "title": "开启风扇降温", "status": "in_progress"},
    {"id": "2", "title": "开启灌溉系统", "status": "pending"}
  ]
}
</TASK_LIST>'''
    doc.add_paragraph(task_list, style='No Spacing')
    
    doc.add_paragraph('交互逻辑：分步执行，每次只执行一个任务')
    
    # 版本 5
    doc.add_heading('版本 5：工具调用', level=1)
    doc.add_paragraph('需求：让 AI 直接控制设备，而不是只给建议')
    doc.add_paragraph('工具定义：')
    doc.add_paragraph('toggle_relay: 控制设备开关（device: 设备编号，value: 1=开启/0=关闭）')
    tool_example = '{"action": "toggle_relay", "device": 6, "value": 1}'
    doc.add_paragraph(tool_example, style='No Spacing')
    
    doc.add_paragraph('执行协议：按部就班、指令一致、状态核对、等待反馈')
    
    # 版本 6
    doc.add_heading('版本 6：完整 Agent', level=1)
    doc.add_paragraph('需求：让 AI 成为主动管理的智能管家')
    doc.add_paragraph('角色设定：莓小智，大棚智能管家')
    doc.add_paragraph('工作模式：')
    modes = [
        '感知 (Perceive)：实时监控传感器数据',
        '思考 (Think)：基于数据推理识别问题',
        '决策 (Decide)：制定解决方案',
        '行动 (Act)：生成指令并监控执行'
    ]
    for mode in modes:
        doc.add_paragraph(mode, style='List Number')
    
    doc.add_paragraph('核心原则：严禁编造数据')
    
    # 版本对比
    doc.add_heading('版本对比', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ['版本', '能力', '解决问题']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        hdr_cells[i].paragraphs[0].runs[0].bold = True
    
    data = [
        ('v1', '通用问答', '基础咨询'),
        ('v2', '实时数据感知', '数据针对性'),
        ('v3', '专业诊断', '问题识别'),
        ('v4', '分步执行', '执行计划'),
        ('v5', '自动执行', '自动化控制'),
        ('v6', '智能体闭环', '主动管理')
    ]
    
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = cell_data
    
    doc.add_page_break()
    
    # 演示流程
    doc.add_heading('演示流程（3-5 分钟）', level=1)
    steps = [
        ('v1 基础版（30 秒）', '展示基础提示词'),
        ('v2 数据注入（1 分钟）', '添加设备数据，强调实时注入'),
        ('v3 诊断逻辑（1 分钟）', '添加专家知识库'),
        ('v4 任务规划（1 分钟）', '添加 Task List 格式'),
        ('v5 工具调用（1 分钟）', '展示 JSON 指令'),
        ('v6 完整 Agent（30 秒）', '总结智能体闭环')
    ]
    
    for i, (title, content) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.add_run(f'{i}. {title}').bold = True
        doc.add_paragraph(content)
    
    # 关键要点
    doc.add_heading('关键要点', level=1)
    points = [
        '需求驱动：每个版本解决前一个版本的问题',
        '渐进式优化：逐步迭代，不追求一步到位',
        '数据注入：在 AI 推理前注入，每次询问都刷新',
        '工具调用：标准化 JSON 格式，状态核对避免重复'
    ]
    for point in points:
        doc.add_paragraph(point, style='List Bullet')
    
    # 页脚信息
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run('莓界智慧农业平台 | Next.js + TypeScript + MCP + AI Agent')
    footer_run.font.size = Pt(10)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)
    
    # 保存文档
    doc.save('prompt-engineering-simple.docx')
    print('精简版 Word 文档已生成：prompt-engineering-simple.docx')

if __name__ == '__main__':
    create_word_document()
