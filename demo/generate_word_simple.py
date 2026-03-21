"""
简化版：生成 Word 文档
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_word_document():
    doc = Document()
    
    # 设置中文字体
    doc.styles['Normal'].font.name = u'微软雅黑'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'微软雅黑')
    
    # 标题
    heading = doc.add_heading('需求驱动式提示词工程编写', level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 副标题
    subtitle = doc.add_paragraph('智慧农业平台 AI 助手 - Agent 开发工作流程')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()
    
    # 目录
    doc.add_heading('目录', level=1)
    toc = [
        '版本 1：基础问答助手',
        '版本 2：数据注入',
        '版本 3：诊断逻辑',
        '版本 4：任务规划',
        '版本 5：工具调用',
        '版本 6：完整 Agent',
        '版本对比总结',
        '现场演示流程',
        '关键要点'
    ]
    for item in toc:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_page_break()
    
    # 版本 1
    doc.add_heading('版本 1：基础问答助手', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户场景：农场主想咨询大棚种植问题')
    doc.add_paragraph('"我需要一个 AI 助手，能够回答用户关于大棚种植的基本问题。"')
    
    doc.add_heading('提示词编写', level=2)
    code1 = 'export const BASIC_PROMPT = `\n你是一个智慧农业助手的 AI。\n请回答用户关于大棚种植的问题。\n`;'
    doc.add_paragraph(code1, style='No Spacing')
    
    doc.add_heading('能做什么', level=2)
    doc.add_paragraph('回答通用种植知识', style='List Bullet')
    doc.add_paragraph('解释农业术语', style='List Bullet')
    doc.add_paragraph('提供基础建议', style='List Bullet')
    
    doc.add_heading('存在的问题', level=2)
    p1 = doc.add_paragraph()
    p1.add_run('问题 1：无法获取实时数据').bold = True
    doc.add_paragraph('用户问："现在大棚温度怎么样？"\nAI 答："抱歉，我无法获取实时数据。"')
    
    p2 = doc.add_paragraph()
    p2.add_run('问题 2：建议缺乏针对性').bold = True
    doc.add_paragraph('用户问："我该怎么调节温度？"\nAI 答："可以通过通风、遮阳等方式调节温度。"')
    
    p3 = doc.add_paragraph()
    p3.add_run('问题 3：无法执行实际操作').bold = True
    doc.add_paragraph('用户说："帮我打开风扇"\nAI 答："建议您手动开启风扇..."')
    
    doc.add_heading('下一步优化方向', level=2)
    doc.add_paragraph('我们需要让 AI 能够获取实时数据，这样才能提供精准的建议。')
    
    doc.add_page_break()
    
    # 版本 2
    doc.add_heading('版本 2：数据注入', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户需求：')
    doc.add_paragraph('"我希望 AI 能够知道当前大棚的实际情况，而不是泛泛而谈。"')
    
    doc.add_paragraph('技术需求：')
    doc.add_paragraph('"在后端中间件中，我们需要在用户每次询问助手时，注入数据提示词。"')
    
    doc.add_heading('提示词编写', level=2)
    code2 = '''export const DATA_ENHANCED_PROMPT = `
你是一个智慧农业助手的 AI。

## 当前设备数据（实时注入）
- 室内温度：32.5°C
- 室内湿度：55%
- 光照强度：12000 lux
- CO₂浓度：380 ppm
- 土壤温度：28.2°C
- 土壤湿度：35%

## 重要原则
请基于上述真实数据回答用户的问题，严禁编造数据。
`;'''
    doc.add_paragraph(code2, style='No Spacing')
    
    doc.add_heading('改进效果', level=2)
    p1 = doc.add_paragraph()
    p1.add_run('效果 1：可以回答实时数据问题').bold = True
    doc.add_paragraph('用户问："现在大棚温度怎么样？"\nAI 答："当前室内温度是 32.5°C。"')
    
    p2 = doc.add_paragraph()
    p2.add_run('效果 2：建议更有针对性').bold = True
    doc.add_paragraph('用户问："我该怎么调节温度？"\nAI 答："当前温度 32.5°C，偏高。建议开启风扇降温。"')
    
    doc.add_heading('仍存在问题', level=2)
    doc.add_paragraph('问题 1：缺乏专业诊断')
    doc.add_paragraph('问题 2：没有给出具体建议')
    
    doc.add_page_break()
    
    # 版本 3
    doc.add_heading('版本 3：诊断逻辑', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户需求：')
    doc.add_paragraph('"我希望 AI 不仅能告诉我数据，还能告诉我这些数据是否正常，有问题该怎么解决。"')
    
    doc.add_heading('提示词编写', level=2)
    doc.add_paragraph('添加草莓种植最佳环境参数（农业专家知识库）：')
    params = [
        '温度：18-28°C（白天），12-15°C（夜间）',
        '湿度：60-75%',
        '光照：15000-30000 lux',
        'CO₂：400-800 ppm（白天）',
        '土壤温度：15-25°C',
        '土壤湿度：40-60%'
    ]
    for param in params:
        doc.add_paragraph(param, style='List Bullet')
    
    doc.add_heading('改进效果', level=2)
    doc.add_paragraph('效果 1：专业的数据对比', style='List Bullet')
    doc.add_paragraph('效果 2：问题识别与优先级', style='List Bullet')
    doc.add_paragraph('效果 3：具体的调节建议', style='List Bullet')
    
    doc.add_page_break()
    
    # 版本 4
    doc.add_heading('版本 4：任务规划', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户需求：')
    doc.add_paragraph('"我希望 AI 能够帮我制定一个执行计划，告诉我先做什么、后做什么。"')
    
    doc.add_heading('任务规划能力', level=2)
    doc.add_paragraph('当用户提出复杂目标或发现多个问题时，请生成结构化的任务清单。')
    
    task_list_example = '''<TASK_LIST>
{
  "title": "大棚环境优化计划",
  "status": "in_progress",
  "tasks": [
    {"id": "1", "title": "开启风扇降温", "status": "in_progress"},
    {"id": "2", "title": "开启灌溉系统", "status": "pending"}
  ]
}
</TASK_LIST>'''
    doc.add_paragraph(task_list_example, style='No Spacing')
    
    doc.add_heading('交互逻辑', level=2)
    logic = [
        '初次规划：分析完环境后，生成完整的 Task List',
        '分步执行：每次只执行一个任务',
        '状态更新：每执行完一步，重新输出更新后的 Task List',
        '完成任务：当所有步骤都完成后，输出状态为 completed'
    ]
    for i, item in enumerate(logic, 1):
        doc.add_paragraph(f'{i}. {item}', style='List Number')
    
    doc.add_page_break()
    
    # 版本 5
    doc.add_heading('版本 5：工具调用', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户需求：')
    doc.add_paragraph('"我希望 AI 能够直接帮我控制设备，而不是只告诉我该怎么做。"')
    
    doc.add_heading('工具调用能力', level=2)
    doc.add_paragraph('工具 1：toggle_relay（控制继电器开关）')
    doc.add_paragraph('描述：控制大棚设备的开关（水泵、风扇、补光灯等）')
    doc.add_paragraph('参数：device（设备编号）, value（1=开启，0=关闭）')
    
    tool_example = '{"action": "toggle_relay", "device": 6, "value": 1}'
    doc.add_paragraph(tool_example, style='No Spacing')
    
    doc.add_heading('执行协议', level=2)
    protocols = [
        '按部就班：每次只执行一个任务',
        '指令一致：JSON 指令必须与任务描述严格对应',
        '状态核对：执行前检查当前设备状态',
        '等待反馈：执行后等待系统反馈'
    ]
    for i, protocol in enumerate(protocols, 1):
        doc.add_paragraph(f'{i}. {protocol}', style='List Number')
    
    doc.add_page_break()
    
    # 版本 6
    doc.add_heading('版本 6：完整 Agent', level=1)
    
    doc.add_heading('需求来源', level=2)
    doc.add_paragraph('用户需求：')
    doc.add_paragraph('"我希望 AI 能够像专业的农业管家一样，主动帮我管理大棚。"')
    
    doc.add_heading('角色设定', level=2)
    doc.add_paragraph('名字：莓小智')
    doc.add_paragraph('身份：大棚智能管家 / 农业智能体 (Agent)')
    doc.add_paragraph('性格：主动、严谨、负责、贴心')
    
    doc.add_heading('工作模式', level=2)
    modes = [
        '感知 (Perceive)：实时监控传感器数据和天气变化',
        '思考 (Think)：基于数据和农业知识库进行推理',
        '决策 (Decide)：制定分步执行的解决方案',
        '行动 (Act)：生成控制指令，并监控执行结果'
    ]
    for i, mode in enumerate(modes, 1):
        doc.add_paragraph(f'{i}. {mode}', style='List Number')
    
    doc.add_heading('核心原则', level=2)
    doc.add_paragraph('严禁编造数据：只能根据提供的真实数据进行回答')
    
    doc.add_page_break()
    
    # 版本对比
    doc.add_heading('版本对比总结', level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ['版本', '能力', '解决问题', '仍存在问题']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        hdr_cells[i].paragraphs[0].runs[0].bold = True
    
    data = [
        ('v1 基础版', '通用问答', '基础咨询', '无实时数据'),
        ('v2 数据注入', '实时数据感知', '数据针对性', '缺乏诊断逻辑'),
        ('v3 诊断逻辑', '专业诊断', '问题识别', '建议无法执行'),
        ('v4 任务规划', '分步执行', '执行计划', '需要手动操作'),
        ('v5 工具调用', '自动执行', '自动化控制', '缺乏主动性'),
        ('v6 完整 Agent', '智能体闭环', '主动管理', '持续优化中')
    ]
    
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = cell_data
    
    doc.add_page_break()
    
    # 演示流程
    doc.add_heading('现场演示流程', level=1)
    doc.add_paragraph('演示步骤（5-8 分钟）')
    
    steps = [
        ('展示 v1 基础版（30 秒）', '打开 prompt-skill.ts，展示基础提示词'),
        ('现场编写 v2 数据注入（1 分钟）', '添加设备数据部分，强调"实时注入"概念'),
        ('添加 v3 诊断逻辑（1.5 分钟）', '添加最佳环境参数和诊断流程'),
        ('实现 v4 任务规划（1.5 分钟）', '添加 Task List 格式'),
        ('集成 v5 工具调用（1.5 分钟）', '添加工具定义，展示 JSON 指令'),
        ('总结 v6 完整 Agent（1 分钟）', '展示完整提示词，强调智能体闭环')
    ]
    
    for i, (title, content) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.add_run(f'{i}. {title}').bold = True
        doc.add_paragraph(content)
    
    doc.add_page_break()
    
    # 关键要点
    doc.add_heading('关键要点', level=1)
    
    doc.add_heading('需求驱动的核心思想', level=2)
    points1 = [
        '从问题出发：每个版本都是为了解决前一个版本的问题',
        '渐进式优化：不追求一步到位，逐步迭代',
        '用户价值导向：每次升级都带来实际的用户价值提升'
    ]
    for point in points1:
        doc.add_paragraph(point, style='List Bullet')
    
    doc.add_heading('提示词工程的最佳实践', level=2)
    points2 = [
        '结构化：使用清晰的章节和格式',
        '具体化：避免模糊的指令，给出具体示例',
        '约束性：明确告诉 AI 什么能做、什么不能做',
        '可测试：每次修改后都要测试效果'
    ]
    for point in points2:
        doc.add_paragraph(point, style='List Bullet')
    
    doc.add_heading('技术实现的关键点', level=2)
    points3 = [
        '数据注入时机：在 AI 推理之前注入',
        '数据刷新机制：每次用户询问都刷新',
        '工具调用协议：标准化的 JSON 格式',
        '状态管理：任务状态的持久化和更新'
    ]
    for point in points3:
        doc.add_paragraph(point, style='List Bullet')
    
    # 页脚信息
    doc.add_page_break()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run('文档版本：v1.0 | 创建日期：2026 年 3 月 20 日 | 项目名称：莓界智慧农业平台')
    footer_run.font.size = Pt(10)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)
    
    # 保存文档
    doc.save('prompt-engineering-step-by-step.docx')
    print('Word 文档已生成：prompt-engineering-step-by-step.docx')

if __name__ == '__main__':
    create_word_document()
